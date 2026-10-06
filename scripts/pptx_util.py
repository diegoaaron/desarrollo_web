# -*- coding: utf-8 -*-
"""Utilidades de maquetado para la presentacion, con la identidad visual del grupo."""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

# ------------------------------------------------------------------- paleta
DARK = RGBColor(0x3D, 0x19, 0x27)
PRIMARY = RGBColor(0x6D, 0x2E, 0x46)
ACCENT = RGBColor(0xA2, 0x67, 0x69)
ROSE = RGBColor(0xD9, 0xAF, 0xAF)
CREAM = RGBColor(0xEC, 0xE2, 0xD0)
MUTED_LIGHT = RGBColor(0xD8, 0xC6, 0xC0)
CARD = RGBColor(0xFB, 0xF7, 0xF2)
BLUE = RGBColor(0x3C, 0x6E, 0x8F)
GREEN = RGBColor(0x4E, 0x7A, 0x5A)
GOLD = RGBColor(0x8A, 0x6D, 0x3B)
TEXT = RGBColor(0x2B, 0x20, 0x20)
MUTED = RGBColor(0x6E, 0x5B, 0x5B)
LINE = RGBColor(0xD8, 0xC6, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TITULO = "Cambria"
CUERPO = "Calibri"

ANCHO = 13.333
ALTO = 7.5

FIGURAS = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "clase_entregables", "figuras")


def nueva_presentacion():
    prs = Presentation()
    prs.slide_width = Inches(ANCHO)
    prs.slide_height = Inches(ALTO)
    return prs


def _fondo(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def lamina(prs, fondo=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # en blanco
    _fondo(slide, fondo)
    return slide


def caja(slide, x, y, w, h, color=CARD, borde=None, grosor=0.75, radio=0.04,
         forma=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde
        sh.line.width = Pt(grosor)
    sh.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        # radio de esquina proporcional al lado menor
        try:
            sh.adjustments[0] = radio / min(w, h) if min(w, h) else 0.1
        except (IndexError, ZeroDivisionError):
            pass
    sh.text_frame.text = ""
    return sh


def _espaciado(run, valor):
    """Aplica espaciado entre caracteres (en centesimas de punto)."""
    run.font._rPr.set("spc", str(int(valor)))


def texto(slide, x, y, w, h, contenido, size=14, color=TEXT, bold=False,
          font=CUERPO, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          interlineado=1.2, spc=None, italic=False, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lineas = contenido.split("\n")
    for i, linea in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = interlineado
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = linea
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
        if spc:
            _espaciado(r, spc)
    return tb


def cabecera(slide, kicker, titulo, color_titulo=PRIMARY):
    texto(slide, 0.7, 0.42, 11.9, 0.3, kicker.upper(), size=11, color=ACCENT,
          bold=True, spc=200)
    texto(slide, 0.7, 0.74, 11.9, 0.7, titulo, size=30, color=color_titulo,
          bold=True, font=TITULO)


def chip(slide, x, y, w, h, etiqueta, color=PRIMARY, texto_color=WHITE, size=9):
    caja(slide, x, y, w, h, color=color, radio=0.03)
    texto(slide, x, y, w, h, etiqueta, size=size, color=texto_color, bold=True,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def tarjeta(slide, x, y, w, h, titulo, cuerpo, color=PRIMARY, size_t=14,
            size_c=10.5, etiqueta=None, fondo=CARD):
    caja(slide, x, y, w, h, color=fondo, borde=LINE)
    caja(slide, x, y, 0.07, h, color=color, radio=0.02, forma=MSO_SHAPE.RECTANGLE)
    dy = 0.26
    if etiqueta:
        chip(slide, x + dy, y + dy, 0.95, 0.26, etiqueta, color=color, size=8)
        texto(slide, x + dy + 1.12, y + dy - 0.03, w - dy - 1.3, 0.32, titulo,
              size=size_t, color=color, bold=True, font=TITULO)
    else:
        texto(slide, x + dy, y + dy, w - 2 * dy, 0.34, titulo, size=size_t,
              color=color, bold=True, font=TITULO)
    texto(slide, x + dy, y + dy + 0.44, w - 2 * dy, h - dy - 0.6, cuerpo,
          size=size_c, color=TEXT, interlineado=1.32, space_after=3)


def vinetas(slide, x, y, w, items, size=11.5, color=TEXT, interlineado=1.35,
            bullet_color=ACCENT, sep=0.06):
    """Lista con vinetas dibujadas como puntos."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = interlineado
        p.space_after = Pt(sep * 72)
        rb = p.add_run()
        rb.text = "▪  "
        rb.font.size = Pt(size)
        rb.font.color.rgb = bullet_color
        rb.font.name = CUERPO
        if isinstance(it, tuple):
            cabeza, resto = it
            r1 = p.add_run()
            r1.text = cabeza
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = color
            r1.font.name = CUERPO
            r2 = p.add_run()
            r2.text = resto
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = CUERPO
        else:
            r = p.add_run()
            r.text = it
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = CUERPO
    return tb


def figura(slide, archivo, x=None, y=1.62, ancho_max=12.2, alto_max=5.5):
    """Inserta una figura escalada para caber en el area disponible."""
    from PIL import Image
    ruta = os.path.join(FIGURAS, archivo)
    with Image.open(ruta) as im:
        pw, ph = im.size
    escala = min(ancho_max / pw, alto_max / ph)
    w, h = pw * escala, ph * escala
    if x is None:
        x = (ANCHO - w) / 2
    slide.shapes.add_picture(ruta, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def pie(slide, contenido, color=MUTED):
    texto(slide, 0.7, 6.98, 11.9, 0.3, contenido, size=9.5, color=color,
          italic=True)


def numero(slide, n, total):
    texto(slide, 12.0, 6.98, 0.8, 0.3, f"{n} / {total}", size=9.5, color=MUTED,
          align=PP_ALIGN.RIGHT)
