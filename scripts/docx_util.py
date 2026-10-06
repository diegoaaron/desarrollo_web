# -*- coding: utf-8 -*-
"""Utilidades de formato para construir el informe en Word."""

import os

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

# ------------------------------------------------------------------- paleta
PRIMARY = RGBColor(0x6D, 0x2E, 0x46)
ACCENT = RGBColor(0xA2, 0x67, 0x69)
BLUE = RGBColor(0x3C, 0x6E, 0x8F)
TEXT = RGBColor(0x2B, 0x20, 0x20)
MUTED = RGBColor(0x6E, 0x5B, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_PRIMARY = "6D2E46"
HEX_ACCENT = "A26769"
HEX_BLUE = "3C6E8F"
HEX_CARD = "FBF7F2"
HEX_ZEBRA = "F6F0EC"
HEX_LINE = "D8C6C0"

FIGURAS = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "clase_entregables", "figuras")


# ------------------------------------------------------------------ formato
def _set_font(style, name="Arial"):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def nuevo_documento():
    """Documento con el formato que exige la consigna: Arial 11, interlineado 1.5."""
    doc = Document()

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.54)
        s.left_margin = s.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    _set_font(normal)
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    encabezados = [
        ("Heading 1", 16, PRIMARY, 18, 10),
        ("Heading 2", 13, PRIMARY, 14, 6),
        ("Heading 3", 11.5, BLUE, 10, 4),
        ("Heading 4", 11, ACCENT, 8, 3),
    ]
    for nombre, tam, color, antes, despues in encabezados:
        st = doc.styles[nombre]
        _set_font(st)
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(antes)
        st.paragraph_format.space_after = Pt(despues)
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for nombre in ("List Bullet", "List Number"):
        st = doc.styles[nombre]
        _set_font(st)
        st.font.size = Pt(11)
        st.font.color.rgb = TEXT
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(3)

    return doc


def sombrear(celda, hexcolor):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tc_pr.append(shd)


def _bordes_tabla(tabla, color=HEX_LINE, sz=4):
    tbl_pr = tabla._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _repetir_cabecera(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def parrafo_celda(celda, texto, size=9, bold=False, color=TEXT, align=None,
                  space_after=2):
    p = celda.paragraphs[0]
    p.text = ""
    partes = texto.split("\n")
    for i, parte in enumerate(partes):
        if i:
            p = celda.add_paragraph()
        r = p.add_run(parte)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Arial"
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        if align is not None:
            p.alignment = align


def tabla(doc, cabeceras, filas, anchos=None, size=9, color_cabecera=HEX_PRIMARY,
          zebra=True, alineaciones=None):
    """Crea una tabla con cabecera de color y filas alternadas."""
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _bordes_tabla(t)

    hdr = t.rows[0]
    _repetir_cabecera(hdr)
    for i, c in enumerate(cabeceras):
        sombrear(hdr.cells[i], color_cabecera)
        parrafo_celda(hdr.cells[i], c, size=size, bold=True, color=WHITE)

    for j, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            if zebra and j % 2 == 1:
                sombrear(celdas[i], HEX_ZEBRA)
            al = alineaciones[i] if alineaciones else None
            parrafo_celda(celdas[i], str(valor), size=size, align=al)

    if anchos:
        total = sum(anchos)
        util = doc.sections[-1].page_width - doc.sections[-1].left_margin - \
            doc.sections[-1].right_margin
        for fila in t.rows:
            for i, c in enumerate(fila.cells):
                c.width = int(util * anchos[i] / total)
    doc.add_paragraph()
    return t


def figura(doc, archivo, pie, ancho_pulgadas=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIGURAS, archivo), width=Inches(ancho_pulgadas))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(pie)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = MUTED
    r.font.name = "Arial"


def figura_apaisada(doc, archivo, pie, ancho_pulgadas=9.0):
    """Coloca la figura en su propia pagina horizontal y vuelve a vertical."""
    seccion_horizontal(doc)
    figura(doc, archivo, pie, ancho_pulgadas=ancho_pulgadas)
    seccion_vertical(doc)


def seccion_horizontal(doc):
    """Abre una sección apaisada y devuelve el objeto sección."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    ancho, alto = s.page_width, s.page_height
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = alto, ancho
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.0)
    return s


def seccion_vertical(doc):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    ancho, alto = s.page_width, s.page_height
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = alto, ancho
    s.top_margin = s.bottom_margin = Cm(2.54)
    s.left_margin = s.right_margin = Cm(2.54)
    return s


def indice(doc):
    """Inserta un campo TOC que Word rellena al abrir el documento (F9)."""
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    ph = OxmlElement("w:p")
    pr = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Haga clic aquí y pulse F9 para generar el índice."
    pr.append(t)
    ph.append(pr)
    fld.append(ph)
    r._r.addnext(fld)


def salto_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def parrafo(doc, texto, size=11, bold=False, italic=False, color=TEXT,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(texto)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Arial"
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def vineta(doc, texto, negrita_hasta=None):
    """Viñeta; si negrita_hasta es un separador, la parte previa va en negrita."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if negrita_hasta and negrita_hasta in texto:
        cabeza, resto = texto.split(negrita_hasta, 1)
        r1 = p.add_run(cabeza + negrita_hasta)
        r1.font.bold = True
        r2 = p.add_run(resto)
        for r in (r1, r2):
            r.font.size = Pt(11)
            r.font.name = "Arial"
            r.font.color.rgb = TEXT
    else:
        r = p.add_run(texto)
        r.font.size = Pt(11)
        r.font.name = "Arial"
        r.font.color.rgb = TEXT
    return p


def bloque_codigo(doc, texto, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(texto)
    r.font.name = "Consolas"
    r.font.size = Pt(size)
    r.font.color.rgb = TEXT
    rpr = r._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Consolas")
    return p


def nota(doc, titulo, texto):
    """Recuadro de una sola celda para destacar una idea."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    _bordes_tabla(t, color=HEX_ACCENT, sz=6)
    c = t.rows[0].cells[0]
    sombrear(c, HEX_CARD)
    parrafo_celda(c, titulo.upper(), size=9, bold=True, color=ACCENT, space_after=4)
    p = c.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(10)
    r.font.name = "Arial"
    r.font.color.rgb = TEXT
    p.paragraph_format.line_spacing = 1.3
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    return t
