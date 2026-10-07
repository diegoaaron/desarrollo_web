# -*- coding: utf-8 -*-
"""Genera clase_entregables/avance_2_ver_1.pptx (avance 2: arquitectura de la solución).

La portada es la de la presentación del entregable 1 (presentacion.portada); las demás
láminas presentan los cuatro patrones de arquitectura vistos en clase aplicados a Coral Shop."""

import os

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import esquema
from presentacion import portada
from pptx_util import (
    nueva_presentacion, lamina, caja, texto, cabecera, chip, tarjeta, vinetas, figura, pie, numero,
    DARK, PRIMARY, ACCENT, ROSE, CREAM, MUTED_LIGHT, CARD, BLUE, GREEN, GOLD, TEXT, MUTED, LINE, WHITE,
    TITULO,
)

SALIDA = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "clase_entregables", "avance_2_ver_1.pptx")

C = PP_ALIGN.CENTER
M = MSO_ANCHOR.MIDDLE


def portada_avance(prs):
    portada(prs, lema="Avance 2  ·  Arquitectura de la solución: capas, cliente-servidor, MVC y eventos")


def agenda(prs):
    s = lamina(prs)
    cabecera(s, "Agenda", "Lo que vamos a recorrer")
    bloques = [
        ("1", "Dónde estamos", "Fases terminadas y el caso de los 12 polos", PRIMARY),
        ("2", "Arquitectura en capas", "El restaurante y el recorrido de un pedido", ACCENT),
        ("3", "Cliente-servidor", "Varios clientes, un servidor, una base", BLUE),
        ("4", "MVC", "React como vista, Spring como controlador", GREEN),
        ("5", "Dirigida por eventos", "OrderStatusChanged, la propuesta de la fase 4", GOLD),
        ("6", "Implementación", "Paquetes, reglas, pruebas y próximos pasos", PRIMARY),
    ]
    for i, (n, t, d, color) in enumerate(bloques):
        fila, col = divmod(i, 3)
        x = 0.7 + col * 4.05
        y = 1.95 + fila * 2.35
        caja(s, x, y, 3.75, 1.95, color=CARD, borde=LINE)
        chip(s, x + 0.28, y + 0.26, 0.42, 0.42, n, color=color, size=13)
        texto(s, x + 0.28, y + 0.86, 3.2, 0.34, t, size=15, color=color, bold=True, font=TITULO)
        texto(s, x + 0.28, y + 1.28, 3.2, 0.5, d, size=10.5, color=MUTED, interlineado=1.3)
    pie(s, "Avance 2  ·  Arquitectura de la solución")


def estado(prs):
    s = lamina(prs)
    cabecera(s, "1 · Dónde estamos", "Tres fases terminadas e integradas")
    fases = [
        ("0", "Ordenar la casa", "Terminada", GREEN), ("1", "Modelo de datos", "Terminada", GREEN),
        ("2", "API de compra", "Terminada", GREEN), ("3", "Frontend", "Siguiente", ACCENT),
        ("4", "Eventos", "Diseñada", GOLD), ("5", "Calidad y demo", "Pendiente", MUTED),
        ("6", "Pasarela real", "Pendiente", MUTED),
    ]
    ancho = 11.93 / len(fases)
    for i, (n, t, e, color) in enumerate(fases):
        x = 0.7 + i * ancho
        caja(s, x + 0.04, 1.65, ancho - 0.08, 1.45, color=CARD, borde=LINE)
        chip(s, x + 0.2, 1.82, 0.42, 0.42, n, color=color, size=12)
        texto(s, x + 0.2, 2.36, ancho - 0.4, 0.4, t, size=11.5, color=TEXT, bold=True, interlineado=1.1)
        texto(s, x + 0.2, 2.72, ancho - 0.4, 0.3, e, size=9.5, color=color, italic=True)

    n_tablas, _, _ = esquema.totales()
    cifras = [(str(n_tablas), "tablas en PostgreSQL"), ("13", "módulos del backend"), ("40", "pruebas unitarias"),
              ("64", "controles de punta a punta")]
    for i, (n, d) in enumerate(cifras):
        x = 0.7 + i * 3.0
        texto(s, x, 3.45, 2.8, 0.75, n, size=40, color=PRIMARY, bold=True, font=TITULO)
        texto(s, x, 4.25, 2.8, 0.3, d, size=11, color=MUTED)

    caja(s, 0.7, 4.95, 11.93, 1.75, color=PRIMARY)
    texto(s, 1.0, 5.15, 11.3, 0.25, "EL CASO QUE GUÍA TODO EL AVANCE", size=9.5, color=ROSE, bold=True, spc=150)
    texto(s, 1.0, 5.48, 11.3, 1.1,
          "Un cliente personaliza 12 polos con su logo bordado en el pecho (4 S, 4 M, 4 L), crea el pedido, paga, "
          "y el administrador lo avanza de PAGADO a ENTREGADO. Total verificado: S/ 591.60 + S/ 10.00 de envío.",
          size=13, color=CREAM, interlineado=1.3)


def patrones(prs):
    s = lamina(prs)
    cabecera(s, "Patrones vistos en clase", "Cuatro patrones, un mismo sistema")
    cards = [
        ("1 · CAPAS", "Organiza el servidor",
         "Presentación → Negocio → Datos. Cada módulo tiene controller, service y repository.", PRIMARY),
        ("2 · CLIENTE-SERVIDOR", "Reparte qué corre dónde",
         "React en el navegador, Spring Boot en el servidor y PostgreSQL detrás.", BLUE),
        ("3 · MVC", "Separa datos, pantalla y coordinación",
         "Modelo: records y servicios. Vista: React. Controlador: @RestController.", GREEN),
        ("4 · DIRIGIDO POR EVENTOS", "Reacciona a hechos del negocio",
         "OrderStatusChanged publicado en el bus de Spring. Diseñado para la fase 4.", GOLD),
    ]
    for i, (et, t, d, color) in enumerate(cards):
        fila, col = divmod(i, 2)
        x = 0.7 + col * 6.05
        y = 1.75 + fila * 2.55
        caja(s, x, y, 5.85, 2.3, color=CARD, borde=LINE)
        chip(s, x + 0.3, y + 0.3, 2.9, 0.36, et, color=color, size=10)
        texto(s, x + 0.3, y + 0.85, 5.3, 0.4, t, size=17, color=color, bold=True, font=TITULO)
        texto(s, x + 0.3, y + 1.38, 5.3, 0.8, d, size=12, color=TEXT, interlineado=1.3)
    pie(s, "No compiten: describen el sistema desde niveles distintos y se aplican a la vez.")


def capas_restaurante(prs):
    s = lamina(prs)
    cabecera(s, "2 · Arquitectura en capas", "El mozo no entra a la despensa")
    figura(s, "A2_01_capas_restaurante.png", y=1.5, ancho_max=11.6, alto_max=5.35)
    pie(s, "En el código: ningún controller recibe JdbcTemplate ni escribe SQL; todo el SQL vive en los @Repository.")


def capas_pedido(prs):
    s = lamina(prs)
    cabecera(s, "2 · Arquitectura en capas", "Crear un pedido: tres capas, tres responsabilidades")
    figura(s, "A2_02_capas_pedido.png", x=0.7, y=1.5, ancho_max=7.2, alto_max=5.4)
    vinetas(s, 8.1, 1.75, 4.6, [
        ("Presentación. ", "El DTO valida el formato con @Valid; si falla, 400 sin tocar la base."),
        ("Negocio. ", "Una transacción: producto activo, zona válida para la técnica, diseño propio, escala de "
                      "mayoreo y precios recalculados."),
        ("Datos. ", "Descuento condicional de stock; si no alcanza, rollback completo y 409."),
        ("Modelo compartido. ", "PricedLine, NewOrder y OrderStatus los usan las tres capas."),
    ], size=12, interlineado=1.3, sep=0.12)


def errores(prs):
    s = lamina(prs)
    cabecera(s, "2 · Arquitectura en capas", "Cada capa responde con su propio código")
    cards = [
        ("400", "Presentación", "Formato inválido",
         "Cotización sin líneas · archivo que no es imagen · diseño de más de 5 MB", BLUE),
        ("404", "Negocio", "No existe para ti", "Un cliente pide el detalle de un pedido ajeno", MUTED),
        ("409", "Negocio · Datos", "Conflicto", "«Stock insuficiente para POLO-…-S (S, Blanco): pediste 21»",
         ACCENT),
        ("422", "Negocio", "Regla violada", "«La zona ESPALDA no admite bordado…» · pasar de PAGADO a ENTREGADO",
         PRIMARY),
    ]
    for i, (cod, capa, t, ej, color) in enumerate(cards):
        x = 0.7 + i * 3.02
        caja(s, x, 1.75, 2.82, 4.2, color=CARD, borde=LINE)
        texto(s, x + 0.3, 1.95, 2.3, 0.9, cod, size=44, color=color, bold=True, font=TITULO)
        chip(s, x + 0.3, 2.95, 2.2, 0.32, capa, color=color, size=9)
        texto(s, x + 0.3, 3.5, 2.3, 0.4, t, size=14, color=TEXT, bold=True, font=TITULO)
        texto(s, x + 0.3, 4.0, 2.3, 1.8, ej, size=11, color=MUTED, interlineado=1.3, italic=True)
    caja(s, 0.7, 6.15, 11.93, 0.62, color=PRIMARY)
    texto(s, 1.0, 6.15, 11.3, 0.62, "Un único @RestControllerAdvice devuelve siempre { status, message, errors }: "
          "mensajes reales obtenidos al verificar la API.", size=12, color=CREAM, anchor=M)


def cliente_servidor(prs):
    s = lamina(prs)
    cabecera(s, "3 · Cliente-servidor", "El servidor es la única autoridad")
    figura(s, "A2_03_cliente_servidor.png", y=1.5, ancho_max=11.4, alto_max=5.35)
    pie(s, "El servidor recalcula precios y stock: nunca acepta importes enviados por el navegador.")


def mvc(prs):
    s = lamina(prs)
    cabecera(s, "4 · Modelo-Vista-Controlador", "Los tres roles de clase, con la vista en el navegador")
    figura(s, "A2_04_mvc.png", y=1.5, ancho_max=11.4, alto_max=5.35)
    pie(s, "JavaBeans → Modelo · Servlets → Controlador · JSP → Vista: el controlador ahora responde JSON y React "
           "lo dibuja.")


def eventos(prs):
    s = lamina(prs)
    cabecera(s, "5 · Dirigida por eventos", "Quien publica no conoce a quien escucha")
    figura(s, "A2_05_eventos.png", y=1.5, ancho_max=11.4, alto_max=5.35)
    pie(s, "Aún no implementado: se agrega en la fase 4 con el bus de Spring, sin broker externo.", color=GOLD)


def conviven(prs):
    s = lamina(prs)
    cabecera(s, "Síntesis", "Cómo conviven los cuatro patrones")
    filas = [
        ("Cliente-servidor", "Despliegue", "Reglas y datos protegidos en un solo lugar", BLUE),
        ("MVC", "Interacción", "La misma API sirve a la tienda, al panel y a las pruebas", GREEN),
        ("Capas", "Organización del servidor", "Cambios acotados y errores 400, 422 y 409 bien separados", PRIMARY),
        ("Eventos", "Reacciones del negocio", "Notificar y auditar sin tocar el flujo de pedidos", GOLD),
    ]
    texto(s, 0.95, 1.65, 3.2, 0.3, "PATRÓN", size=9.5, color=MUTED, bold=True, spc=120)
    texto(s, 4.2, 1.65, 3.2, 0.3, "NIVEL", size=9.5, color=MUTED, bold=True, spc=120)
    texto(s, 7.2, 1.65, 5.2, 0.3, "BENEFICIO PARA CORAL SHOP", size=9.5, color=MUTED, bold=True, spc=120)
    for i, (p, n, b, color) in enumerate(filas):
        y = 2.05 + i * 1.15
        caja(s, 0.7, y, 11.93, 0.95, color=CARD, borde=LINE)
        caja(s, 0.7, y, 0.08, 0.95, color=color, radio=0.02, forma=MSO_SHAPE.RECTANGLE)
        texto(s, 0.95, y, 3.1, 0.95, p, size=16, color=color, bold=True, font=TITULO, anchor=M)
        texto(s, 4.2, y, 2.8, 0.95, n, size=12, color=TEXT, anchor=M)
        texto(s, 7.2, y, 5.2, 0.95, b, size=12, color=TEXT, anchor=M, interlineado=1.2)


def implementacion(prs):
    s = lamina(prs)
    cabecera(s, "6 · Implementación", "Paquetes por funcionalidad y por capa")
    caja(s, 0.7, 1.6, 6.3, 5.15, color=DARK)
    texto(s, 1.0, 1.82, 5.8, 4.8,
          "com.coralshop/\n"
          "├─ auth/          config · controller · dto · service\n"
          "├─ catalog/       controller · dto · model · repository · service\n"
          "├─ customization/ controller · dto · model · repository · service\n"
          "├─ design/        controller · dto · model · repository · service\n"
          "├─ pricing/       controller · dto · model · service\n"
          "├─ order/         controller · dto · model · repository · service\n"
          "├─ payment/       controller · dto · model · repository · service\n"
          "├─ address/ · shipping/ · user/ · stats/ · health/\n"
          "└─ common/        dto · exception\n\n"
          "db/migration/     V1 … V5 (Flyway)",
          size=10, color=CREAM, font="Consolas", interlineado=1.35)
    tarjeta(s, 7.25, 1.6, 5.38, 1.6, "Cliente", "Catálogo filtrado, personalizador, diseños, cotización, "
            "direcciones, pedido, pago y «Mis pedidos».", color=BLUE, size_c=11)
    tarjeta(s, 7.25, 3.37, 5.38, 1.6, "Administrador", "Bandeja de pedidos con máquina de estados, productos, "
            "categorías, usuarios y métricas.", color=PRIMARY, size_c=11)
    tarjeta(s, 7.25, 5.14, 5.38, 1.6, "Seguridad", "Sesión, CSRF, roles (un admin también compra) y SQL siempre "
            "con parámetros.", color=GREEN, size_c=11)


def reglas_pruebas(prs):
    s = lamina(prs)
    cabecera(s, "6 · Implementación", "Reglas del negocio, verificadas")
    vinetas(s, 0.7, 1.7, 6.4, [
        ("Precio en el servidor: ", "base + técnica + zonas, con el descuento de la escala."),
        ("Mayoreo: ", "1–2 sin descuento · 3–5 al 5 % · 6–11 al 10 % · 12+ al 15 %."),
        ("Zonas por tipo: ", "el bordado admite menos zonas y tamaños que el estampado."),
        ("Sin sobreventa: ", "UPDATE … WHERE stock >= ? dentro de la transacción."),
        ("Estados: ", "PAGADO solo con pago aprobado; cancelar repone el stock."),
        ("Pago seguro: ", "pedido bloqueado durante el cobro; la tarjeta no se guarda."),
    ], size=12.5, interlineado=1.3, sep=0.14)
    pruebas = [("40 / 40", "pruebas unitarias", "QuantityPolicy, importes, estados, imagen y pasarela", GREEN),
               ("64 / 64", "controles de punta a punta", "flujo completo y cada regla rota contra PostgreSQL", BLUE),
               ("0", "errores 500", "en el registro del servidor durante la verificación", PRIMARY)]
    for i, (n, t, d, color) in enumerate(pruebas):
        y = 1.7 + i * 1.7
        caja(s, 7.6, y, 5.03, 1.5, color=CARD, borde=LINE)
        texto(s, 7.9, y + 0.15, 2.0, 0.7, n, size=28, color=color, bold=True, font=TITULO)
        texto(s, 10.0, y + 0.25, 2.5, 0.35, t, size=12, color=TEXT, bold=True)
        texto(s, 7.9, y + 0.88, 4.5, 0.5, d, size=10.5, color=MUTED, interlineado=1.2)


def proximos(prs):
    s = lamina(prs)
    cabecera(s, "Próximos pasos", "Del API a la tienda completa")
    pasos = [
        ("Fase 3", "Frontend de compra", "Personalizador con vista previa, carrito por diseño, checkout, pago y "
         "«Mis pedidos».", BLUE),
        ("Fase 4", "Evento de dominio", "OrderStatusChanged con notificación al cliente y auditoría.", GOLD),
        ("Fase 5", "Calidad y demo", "Datos ficticios, colección Postman, informe y presentación finales.", GREEN),
        ("Despliegue", "Vercel + backend", "Frontend en Vercel con la variable BACKEND_URL; backend en el servicio "
         "que se elija.", PRIMARY),
    ]
    for i, (et, t, d, color) in enumerate(pasos):
        x = 0.7 + i * 3.02
        caja(s, x, 1.75, 2.82, 3.3, color=CARD, borde=LINE)
        chip(s, x + 0.28, 2.0, 1.4, 0.36, et, color=color, size=10)
        texto(s, x + 0.28, 2.6, 2.3, 0.8, t, size=17, color=color, bold=True, font=TITULO, interlineado=1.1)
        texto(s, x + 0.28, 3.55, 2.3, 2.4, d, size=11.5, color=TEXT, interlineado=1.3)


def cierre(prs):
    s = lamina(prs, DARK)
    caja(s, -1.4, 4.6, 4.6, 4.6, color=PRIMARY, forma=MSO_SHAPE.OVAL)
    caja(s, 11.0, -1.4, 3.8, 3.8, color=ACCENT, forma=MSO_SHAPE.OVAL)
    texto(s, 1.2, 1.0, 9.5, 0.3, "PARA CERRAR", size=11, color=ROSE, bold=True, spc=200)
    texto(s, 1.2, 1.42, 10.5, 0.9, "La arquitectura se ve en el código", size=38, color=CREAM, bold=True,
          font=TITULO)
    puntos = [
        ("Capas comprobables", "Ningún controller toca la base; cada capa responde con su código."),
        ("Servidor como autoridad", "Precios y stock se recalculan siempre en el servidor."),
        ("MVC adaptado", "La vista pasó de JSP a React; el controlador responde JSON."),
        ("Eventos con propósito", "Se agregan en la fase 4 para notificar y auditar."),
    ]
    for i, (t, d) in enumerate(puntos):
        fila, col = divmod(i, 2)
        x = 1.2 + col * 5.75
        y = 2.75 + fila * 1.55
        caja(s, x, y, 5.4, 1.3, color=PRIMARY)
        texto(s, x + 0.3, y + 0.2, 4.8, 0.34, t, size=14, color=CREAM, bold=True, font=TITULO)
        texto(s, x + 0.3, y + 0.66, 4.8, 0.5, d, size=10.5, color=MUTED_LIGHT, interlineado=1.25)
    texto(s, 1.2, 6.35, 10.5, 0.4, "Gracias.  ¿Preguntas?", size=20, color=ROSE, bold=True, font=TITULO)


def main():
    prs = nueva_presentacion()
    laminas = [portada_avance, agenda, estado, patrones, capas_restaurante, capas_pedido, errores,
               cliente_servidor, mvc, eventos, conviven, implementacion, reglas_pruebas, proximos, cierre]
    for f in laminas:
        f(prs)
    total = len(prs.slides._sldIdLst)
    for i, slide in enumerate(prs.slides, start=1):
        if i in (1, total):
            continue
        numero(slide, i, total)
    prs.save(SALIDA)
    print(f"Presentación del avance 2 generada ({total} láminas):", SALIDA)


if __name__ == "__main__":
    main()
